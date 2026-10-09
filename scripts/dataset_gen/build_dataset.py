import os
import sys
import subprocess
import pypdfium2 as pdfium
import docx
from PIL import Image

# Reconfigure stdout for utf-8 if supported
if hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass

# Add workspace root to sys.path
WORKSPACE_ROOT = r"c:\Users\Windows 10 Lite\Downloads\MUNJOR\4TH YR\DEV\LATEST CLONE\v14\2nd-repo-for-hrms-backend-"
if WORKSPACE_ROOT not in sys.path:
    sys.path.insert(0, WORKSPACE_ROOT)

from scripts.dataset_gen.common import safe_path

from scripts.dataset_gen.marcus_builder import (
    build_marcus_resume1_pdf,
    build_marcus_resume2_docx,
    build_marcus_resume3_png,
    build_marcus_resume4_jpg,
    build_marcus_resume5_blurred_png,
    build_marcus_resume6_blurred_jpg,
    build_marcus_doc1_coe_current,
    build_marcus_doc2_coe_previous,
    build_marcus_doc3_diploma,
    build_marcus_doc4_cert_foodsafety,
    build_marcus_doc5_cert_service,
    build_marcus_doc6_incomplete_cert,
    build_marcus_ground_truth_txt
)

from scripts.dataset_gen.patricia_builder import (
    build_patricia_resume1_pdf,
    build_patricia_resume2_docx,
    build_patricia_resume3_png,
    build_patricia_resume4_jpg,
    build_patricia_resume5_blurred_png,
    build_patricia_resume6_blurred_jpg,
    build_patricia_doc1_coe_current,
    build_patricia_doc2_coe_previous,
    build_patricia_doc3_diploma,
    build_patricia_doc4_cert_coldchain,
    build_patricia_doc5_cert_warehouse,
    build_patricia_doc6_incomplete_cert,
    build_patricia_ground_truth_txt
)

from scripts.dataset_gen.jerome_builder import (
    build_jerome_resume1_pdf,
    build_jerome_resume2_docx,
    build_jerome_resume3_png,
    build_jerome_resume4_jpg,
    build_jerome_resume5_blurred_png,
    build_jerome_resume6_blurred_jpg,
    build_jerome_doc1_coe_current,
    build_jerome_doc2_coe_previous,
    build_jerome_doc3_diploma,
    build_jerome_doc4_cert_driving,
    build_jerome_doc5_cert_dispatch,
    build_jerome_doc6_incomplete_cert,
    build_jerome_ground_truth_txt
)

def run():
    dataset_name = "Hotel_Restaurant_Hospitality_Verification_Dataset_4"
    dataset_dir = os.path.join(WORKSPACE_ROOT, dataset_name)
    os.makedirs(safe_path(dataset_dir), exist_ok=True)

    print("=" * 70)
    print("STARTING DATASET GENERATION:", dataset_name)
    print("=" * 70)

    # ------------------------------------------------------------------
    # 1. CANDIDATE 1: MARCUS ELIJAH NAVARRO
    # ------------------------------------------------------------------
    marcus_folder_name = "Marcus Elijah Navarro - Hospitality Resume and Documents"
    marcus_dir = os.path.join(dataset_dir, marcus_folder_name)
    os.makedirs(safe_path(marcus_dir), exist_ok=True)
    print(f"\n[1/3] Generating Marcus Elijah Navarro files in: {marcus_folder_name}")

    build_marcus_resume1_pdf(os.path.join(marcus_dir, "Marcus_Elijah_Navarro_Banquet_Operations_Supervisor.pdf"))
    build_marcus_resume2_docx(os.path.join(marcus_dir, "Marcus_Elijah_Navarro_Hotel_Events_Service_Coordinator.docx"))
    build_marcus_resume3_png(os.path.join(marcus_dir, "Marcus_Elijah_Navarro_Banquet_Team_Leader.png"))
    build_marcus_resume4_jpg(os.path.join(marcus_dir, "Marcus_Elijah_Navarro_Catering_Service_Supervisor.jpg"))
    build_marcus_resume5_blurred_png(os.path.join(marcus_dir, "Marcus_Elijah_Navarro_Function_Operations_Officer_Blurred.png"))
    build_marcus_resume6_blurred_jpg(os.path.join(marcus_dir, "Marcus_Elijah_Navarro_Hotel_Events_Operations_Coordinator_Blurred.jpg"))

    build_marcus_doc1_coe_current(os.path.join(marcus_dir, "Marcus_Elijah_Navarro_COE_01.pdf"))
    build_marcus_doc2_coe_previous(os.path.join(marcus_dir, "Marcus_Elijah_Navarro_COE_02.pdf"))
    build_marcus_doc3_diploma(os.path.join(marcus_dir, "Marcus_Elijah_Navarro_Diploma.pdf"))
    build_marcus_doc4_cert_foodsafety(os.path.join(marcus_dir, "Marcus_Elijah_Navarro_Training_Certificate_01.pdf"))
    build_marcus_doc5_cert_service(os.path.join(marcus_dir, "Marcus_Elijah_Navarro_Training_Certificate_02.pdf"))
    build_marcus_doc6_incomplete_cert(os.path.join(marcus_dir, "Marcus_Elijah_Navarro_Professional_Certification.pdf"))

    build_marcus_ground_truth_txt(os.path.join(marcus_dir, "Actual Info in the Resume of Marcus Elijah Navarro.txt"))
    print("[OK] Marcus Elijah Navarro complete (13 files).")

    # ------------------------------------------------------------------
    # 2. CANDIDATE 2: PATRICIA ELAINE RAMOS
    # ------------------------------------------------------------------
    patricia_folder_name = "Patricia Elaine Ramos - Hospitality Resume and Documents"
    patricia_dir = os.path.join(dataset_dir, patricia_folder_name)
    os.makedirs(safe_path(patricia_dir), exist_ok=True)
    print(f"\n[2/3] Generating Patricia Elaine Ramos files in: {patricia_folder_name}")

    build_patricia_resume1_pdf(os.path.join(patricia_dir, "Patricia_Elaine_Ramos_Hotel_Purchasing_Supervisor.pdf"))
    build_patricia_resume2_docx(os.path.join(patricia_dir, "Patricia_Elaine_Ramos_Hospitality_Procurement_Coordinator.docx"))
    build_patricia_resume3_png(os.path.join(patricia_dir, "Patricia_Elaine_Ramos_Food_and_Beverage_Purchasing_Officer.png"))
    build_patricia_resume4_jpg(os.path.join(patricia_dir, "Patricia_Elaine_Ramos_Restaurant_Supply_Chain_Coordinator.jpg"))
    build_patricia_resume5_blurred_png(os.path.join(patricia_dir, "Patricia_Elaine_Ramos_Hospitality_Inventory_and_Procurement_Team_Leader_Blurred.png"))
    build_patricia_resume6_blurred_jpg(os.path.join(patricia_dir, "Patricia_Elaine_Ramos_Hotel_Food_Supply_Operations_Coordinator_Blurred.jpg"))

    build_patricia_doc1_coe_current(os.path.join(patricia_dir, "Patricia_Elaine_Ramos_COE_01.pdf"))
    build_patricia_doc2_coe_previous(os.path.join(patricia_dir, "Patricia_Elaine_Ramos_COE_02.pdf"))
    build_patricia_doc3_diploma(os.path.join(patricia_dir, "Patricia_Elaine_Ramos_Diploma.pdf"))
    build_patricia_doc4_cert_coldchain(os.path.join(patricia_dir, "Patricia_Elaine_Ramos_Training_Certificate_01.pdf"))
    build_patricia_doc5_cert_warehouse(os.path.join(patricia_dir, "Patricia_Elaine_Ramos_Training_Certificate_02.pdf"))
    build_patricia_doc6_incomplete_cert(os.path.join(patricia_dir, "Patricia_Elaine_Ramos_Professional_Certification.pdf"))

    build_patricia_ground_truth_txt(os.path.join(patricia_dir, "Actual Info in the Resume of Patricia Elaine Ramos.txt"))
    print("[OK] Patricia Elaine Ramos complete (13 files).")

    # ------------------------------------------------------------------
    # 3. CANDIDATE 3: JEROME VINCENT ALONZO
    # ------------------------------------------------------------------
    jerome_folder_name = "Jerome Vincent Alonzo - Hospitality Resume and Documents"
    jerome_dir = os.path.join(dataset_dir, jerome_folder_name)
    os.makedirs(safe_path(jerome_dir), exist_ok=True)
    print(f"\n[3/3] Generating Jerome Vincent Alonzo files in: {jerome_folder_name}")

    build_jerome_resume1_pdf(os.path.join(jerome_dir, "Jerome_Vincent_Alonzo_Hotel_Transportation_Services_Supervisor.pdf"))
    build_jerome_resume2_docx(os.path.join(jerome_dir, "Jerome_Vincent_Alonzo_Resort_Guest_Mobility_Coordinator.docx"))
    build_jerome_resume3_png(os.path.join(jerome_dir, "Jerome_Vincent_Alonzo_Hotel_Airport_Transfer_Coordinator.png"))
    build_jerome_resume4_jpg(os.path.join(jerome_dir, "Jerome_Vincent_Alonzo_Hospitality_Shuttle_Operations_Officer.jpg"))
    build_jerome_resume5_blurred_png(os.path.join(jerome_dir, "Jerome_Vincent_Alonzo_Guest_Transport_Team_Leader_Blurred.png"))
    build_jerome_resume6_blurred_jpg(os.path.join(jerome_dir, "Jerome_Vincent_Alonzo_Hotel_Arrival_Services_Coordinator_Blurred.jpg"))

    build_jerome_doc1_coe_current(os.path.join(jerome_dir, "Jerome_Vincent_Alonzo_COE_01.pdf"))
    build_jerome_doc2_coe_previous(os.path.join(jerome_dir, "Jerome_Vincent_Alonzo_COE_02.pdf"))
    build_jerome_doc3_diploma(os.path.join(jerome_dir, "Jerome_Vincent_Alonzo_Diploma.pdf"))
    build_jerome_doc4_cert_driving(os.path.join(jerome_dir, "Jerome_Vincent_Alonzo_Training_Certificate_01.pdf"))
    build_jerome_doc5_cert_dispatch(os.path.join(jerome_dir, "Jerome_Vincent_Alonzo_Training_Certificate_02.pdf"))
    build_jerome_doc6_incomplete_cert(os.path.join(jerome_dir, "Jerome_Vincent_Alonzo_Professional_Certification.pdf"))

    build_jerome_ground_truth_txt(os.path.join(jerome_dir, "Actual Info in the Resume of Jerome Vincent Alonzo.txt"))
    print("[OK] Jerome Vincent Alonzo complete (13 files).")

    # ------------------------------------------------------------------
    # 4. COMPREHENSIVE QUALITY CONTROL & VERIFICATION
    # ------------------------------------------------------------------
    print("\n" + "=" * 70)
    print("PERFORMING QUALITY CONTROL & VERIFICATION ON ALL GENERATED FILES")
    print("=" * 70)

    total_files = 0
    candidate_folders = [marcus_dir, patricia_dir, jerome_dir]

    for c_dir in candidate_folders:
        folder_basename = os.path.basename(c_dir)
        files = os.listdir(safe_path(c_dir))
        print(f"\nVerifying Folder: {folder_basename} (File Count: {len(files)})")
        if len(files) != 13:
            raise ValueError(f"Expected 13 files in {folder_basename}, found {len(files)}")

        resumes = [f for f in files if any(f.endswith(ext) for ext in [".docx", ".png", ".jpg"]) or (f.endswith(".pdf") and not ("COE" in f or "Diploma" in f or "Training" in f or "Professional" in f))]
        docs = [f for f in files if any(k in f for k in ["COE", "Diploma", "Training", "Professional"])]
        txts = [f for f in files if f.endswith(".txt")]

        print(f"  - Resumes ({len(resumes)}): {resumes}")
        print(f"  - Supporting Docs ({len(docs)}): {docs}")
        print(f"  - TXT Ground Truth ({len(txts)}): {txts}")

        if len(resumes) != 6:
            raise ValueError(f"Expected 6 resumes, got {len(resumes)}")
        if len(docs) != 6:
            raise ValueError(f"Expected 6 supporting documents, got {len(docs)}")
        if len(txts) != 1:
            raise ValueError(f"Expected 1 TXT ground truth, got {len(txts)}")

        for fname in files:
            fpath = os.path.join(c_dir, fname)
            fsize = os.path.getsize(safe_path(fpath))
            if fsize == 0:
                raise ValueError(f"File {fname} is 0 bytes (corrupted / empty)!")

            # Test opening each file type
            if fname.endswith(".pdf"):
                doc_pdf = pdfium.PdfDocument(safe_path(fpath))
                assert len(doc_pdf) >= 1
                doc_pdf.close()
            elif fname.endswith(".docx"):
                doc_w = docx.Document(safe_path(fpath))
                assert len(doc_w.paragraphs) > 0
            elif fname.endswith((".png", ".jpg")):
                im = Image.open(safe_path(fpath))
                im.verify()
            elif fname.endswith(".txt"):
                with open(safe_path(fpath), "r", encoding="utf-8") as tf:
                    t_content = tf.read()
                    assert len(t_content) > 1000
                    assert "PART 1" in t_content and "PART 2" in t_content and "PART 3" in t_content and "PART 4" in t_content

            total_files += 1

    print(f"\n[OK] Successfully verified all {total_files} files across 3 candidate folders!")
    assert total_files == 39

    # ------------------------------------------------------------------
    # 5. PACKAGING INTO RAR ARCHIVE
    # ------------------------------------------------------------------
    rar_exe = r"C:\Program Files\WinRAR\Rar.exe"
    archive_path = os.path.join(WORKSPACE_ROOT, f"{dataset_name}.rar")
    if os.path.exists(safe_path(archive_path)):
        os.remove(safe_path(archive_path))

    print(f"\nPackaging into RAR archive: {archive_path}")
    cmd = [
        rar_exe,
        "a",
        "-r",
        f"{dataset_name}.rar",
        dataset_name
    ]
    result = subprocess.run(cmd, cwd=WORKSPACE_ROOT, capture_output=True, text=True)
    if result.returncode != 0:
        print("Rar output:", result.stdout)
        print("Rar error:", result.stderr)
        raise RuntimeError(f"Rar.exe failed with exit code {result.returncode}")

    print("[OK] Archive created successfully.")

    # Test archive integrity
    test_cmd = [rar_exe, "t", f"{dataset_name}.rar"]
    test_res = subprocess.run(test_cmd, cwd=WORKSPACE_ROOT, capture_output=True, text=True)
    if test_res.returncode != 0:
        raise RuntimeError("Rar archive integrity test failed!")
    print("[OK] Archive integrity test: PASSED (All files OK).")

    # List archive bare contents
    list_cmd = [rar_exe, "lb", f"{dataset_name}.rar"]
    list_res = subprocess.run(list_cmd, cwd=WORKSPACE_ROOT, capture_output=True, text=True)
    archived_files = [line.strip() for line in list_res.stdout.splitlines() if line.strip() and not line.strip().endswith("\\") and not line.strip().endswith("/")]
    print(f"[OK] Archive contains {len(archived_files)} total files.")

    print("\n" + "=" * 70)
    print("DATASET GENERATION AND PACKAGING COMPLETE!")
    print(f"Final Archive: {archive_path} (Size: {os.path.getsize(safe_path(archive_path)):,} bytes)")
    print("=" * 70)

if __name__ == "__main__":
    run()
