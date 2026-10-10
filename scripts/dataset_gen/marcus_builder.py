import os
import sys
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.pdfgen import canvas
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from scripts.dataset_gen.common import (
    F_SANS, F_SANS_BOLD, F_SANS_ITALIC,
    F_SERIF, F_SERIF_BOLD, F_SERIF_ITALIC,
    F_CALIBRI, F_CALIBRI_BOLD, F_CALIBRI_ITALIC,
    F_ARIAL, F_ARIAL_BOLD, F_ARIAL_ITALIC,
    F_TIMES, F_TIMES_BOLD, F_TIMES_ITALIC,
    draw_wrapped, pdf_to_image, apply_uniform_blur, apply_torn_scan_defect, safe_path
)

# ----------------------------------------------------------------------
# MARCUS ELIJAH NAVARRO MASTER DATA
# ----------------------------------------------------------------------
MARCUS_PROFILE = {
    "name": "Marcus Elijah Navarro",
    "phone": "+63 917 552 8491",
    "email": "marcus.navarro.hospitality@outlook.ph",
    "address": "Domestic Road, Pasay City, 1301 Metro Manila, Philippines",
    "linkedin": "linkedin.com/in/marcus-navarro-hospitality",
    "education": {
        "degree": "Bachelor of Science in Hospitality Management",
        "institution": "Pamantasan ng Lungsod ng Maynila (PLM) – Intramuros, Manila",
        "year": "2018"
    },
    "experience": [
        {
            "company": "Grand Crest Hotel & Suites",
            "location": "Pasay City, Metro Manila",
            "position": "Banquet Operations Supervisor",
            "duration": "June 2022 – Present",
            "bullets": [
                "Supervise banquet operations for 3 grand ballrooms and 5 multi-function rooms catering up to 600 attendees per event.",
                "Direct pre-shift briefings and manage a team of 18 full-time banquet servers, captains, and on-call catering staff.",
                "Coordinate directly with culinary and sales teams to review BEO specifications, ensuring 99% on-time food rollout.",
                "Reduced banquet equipment loss and chinaware breakage by 18% through standardized end-of-shift inventory protocols."
            ]
        },
        {
            "company": "Harborview Hotel & Convention Center",
            "location": "Manila Bay, Manila",
            "position": "Banquet Team Leader / Captain",
            "duration": "January 2020 – May 2022",
            "bullets": [
                "Led floor teams of 12 banquet attendants during corporate conventions, diplomatic receptions, and wedding galas.",
                "Oversaw banquet table layouts, buffet staging, and silver service standards in compliance with luxury hotel guidelines.",
                "Conducted daily equipment inspections and enforced food temperature logs in coordination with kitchen quality control."
            ]
        },
        {
            "company": "Bayfront Pavilion & Catering Services",
            "location": "Parañaque City, Metro Manila",
            "position": "Banquet Server / Function Attendant",
            "duration": "July 2018 – December 2019",
            "bullets": [
                "Delivered plated dinner service, replenished buffet lines, and executed room turnarounds for corporate and private events.",
                "Maintained high standards of table grooming, glassware polishing, and rapid function hall clearing."
            ]
        }
    ],
    "certifications": [
        {
            "title": "Certified Hospitality Supervisor (CHS)",
            "issuer": "Hospitality Operations Institute / FHCB",
            "date": "November 2021"
        },
        {
            "title": "Food Safety & Sanitation Standards (HACCP Level 2)",
            "issuer": "Hospitality Training Council of the Philippines",
            "date": "August 2022"
        },
        {
            "title": "Customer Service Excellence in Banquets and Events",
            "issuer": "Philippine Hospitality Development Center",
            "date": "March 2020"
        }
    ]
}


# ======================================================================
# RESUME DESIGN A (PDF): Navy Left Sidebar Layout
# Target Role: Banquet Operations Supervisor
# ======================================================================
def build_marcus_resume1_pdf(output_pdf):
    c = canvas.Canvas(safe_path(output_pdf), pagesize=letter)
    w, h = letter

    sb_w = 200
    c.setFillColor(colors.HexColor("#1A2B4C"))
    c.rect(0, 0, sb_w, h, stroke=0, fill=1)

    c.setFillColor(colors.HexColor("#38BDF8"))
    c.rect(sb_w - 4, 0, 4, h, stroke=0, fill=1)

    c.setFillColor(colors.white)
    c.setFont(F_SANS_BOLD, 15)
    c.drawString(18, h - 50, "MARCUS ELIJAH")
    c.drawString(18, h - 68, "NAVARRO")

    c.setFillColor(colors.HexColor("#38BDF8"))
    c.setFont(F_SANS_BOLD, 8.5)
    c.drawString(18, h - 88, "BANQUET OPERATIONS SUPERVISOR")

    curr_y = h - 120
    c.setFillColor(colors.HexColor("#93C5FD"))
    c.setFont(F_SANS_BOLD, 9)
    c.drawString(18, curr_y, "CONTACT INFORMATION")
    curr_y -= 8
    c.setStrokeColor(colors.HexColor("#38BDF8"))
    c.line(18, curr_y, 75, curr_y)
    curr_y -= 14

    contacts = [
        ("PHONE", MARCUS_PROFILE["phone"]),
        ("EMAIL", MARCUS_PROFILE["email"]),
        ("LOCATION", "Domestic Road, Pasay City"),
        ("", "Metro Manila, Philippines"),
        ("LINKEDIN", MARCUS_PROFILE["linkedin"])
    ]
    for label, val in contacts:
        if label:
            c.setFont(F_SANS_BOLD, 7)
            c.setFillColor(colors.HexColor("#60A5FA"))
            c.drawString(18, curr_y, label)
            curr_y -= 9
        c.setFont(F_SANS, 7.5)
        c.setFillColor(colors.white)
        c.drawString(18, curr_y, val)
        curr_y -= 11

    curr_y -= 12
    c.setFillColor(colors.HexColor("#93C5FD"))
    c.setFont(F_SANS_BOLD, 9)
    c.drawString(18, curr_y, "EDUCATION")
    curr_y -= 8
    c.setStrokeColor(colors.HexColor("#38BDF8"))
    c.line(18, curr_y, 75, curr_y)
    curr_y -= 14

    c.setFont(F_SANS_BOLD, 8)
    c.setFillColor(colors.white)
    c.drawString(18, curr_y, "BS in Hospitality Management")
    curr_y -= 10
    c.setFont(F_SANS, 7.5)
    c.setFillColor(colors.HexColor("#CBD5E1"))
    c.drawString(18, curr_y, "Pamantasan ng Lungsod ng Maynila")
    curr_y -= 9
    c.drawString(18, curr_y, "(PLM) – Intramuros, Manila")
    curr_y -= 9
    c.setFont(F_SANS_BOLD, 7.5)
    c.setFillColor(colors.HexColor("#38BDF8"))
    c.drawString(18, curr_y, "Graduated: 2018")

    curr_y -= 22
    c.setFillColor(colors.HexColor("#93C5FD"))
    c.setFont(F_SANS_BOLD, 9)
    c.drawString(18, curr_y, "CERTIFICATIONS")
    curr_y -= 8
    c.setStrokeColor(colors.HexColor("#38BDF8"))
    c.line(18, curr_y, 75, curr_y)
    curr_y -= 14

    certs = [
        ("Certified Hospitality Supervisor", "HOI / FHCB – 2021"),
        ("Food Safety HACCP Level 2", "Hospitality Training Council – 2022"),
        ("Customer Service in Banquets", "PHDC – 2020")
    ]
    for ct, iss in certs:
        c.setFont(F_SANS_BOLD, 7.5)
        c.setFillColor(colors.white)
        c.drawString(18, curr_y, ct)
        curr_y -= 9
        c.setFont(F_SANS, 7)
        c.setFillColor(colors.HexColor("#94A3B8"))
        c.drawString(18, curr_y, iss)
        curr_y -= 13

    mx = sb_w + 18
    mw = w - mx - 20
    my = h - 45

    c.setFillColor(colors.HexColor("#1A2B4C"))
    c.setFont(F_SANS_BOLD, 10.5)
    c.drawString(mx, my, "EXECUTIVE PROFILE")
    my -= 5
    c.setStrokeColor(colors.HexColor("#1A2B4C"))
    c.setLineWidth(1)
    c.line(mx, my, mx + mw, my)
    my -= 14

    summary_text = (
        "Dedicated, results-driven Banquet Operations Supervisor with over 5 years of progressive hospitality "
        "leadership experience in luxury hotel convention centers and high-volume event venues. Proven expertise in "
        "Banquet Event Order (BEO) execution, floor logistics, VIP service protocols, team scheduling, and HACCP "
        "food safety standards. Adept at coordinating functions of up to 600 guests while ensuring seamless service "
        "delivery, minimal equipment breakage, and exceptional guest satisfaction."
    )
    my = draw_wrapped(c, summary_text, mx, my, mw, F_SANS, 8.5, colors.HexColor("#334155"), 11.5)

    my -= 10
    c.setFillColor(colors.HexColor("#1A2B4C"))
    c.setFont(F_SANS_BOLD, 10.5)
    c.drawString(mx, my, "CORE COMPETENCIES & EXPERTISE")
    my -= 5
    c.line(mx, my, mx + mw, my)
    my -= 12

    skills = [
        "• Banquet Event Order (BEO) Floor Planning",
        "• Function Room Turnover & Service Logistics",
        "• Staff Scheduling & Crew Leadership (up to 25)",
        "• VIP Guest Relations & Protocol Handling",
        "• Plated Service, Buffets & Banquet Bars",
        "• Chinaware & Glassware Breakage Mitigation",
        "• Food Safety, Sanitation & HACCP Compliance",
        "• Banquet Billing & Beverage Requisition"
    ]
    col1_w = mw / 2
    for i in range(0, len(skills), 2):
        c.setFont(F_SANS, 8)
        c.setFillColor(colors.HexColor("#1E293B"))
        c.drawString(mx, my, skills[i])
        if i + 1 < len(skills):
            c.drawString(mx + col1_w, my, skills[i + 1])
        my -= 11

    my -= 10
    c.setFillColor(colors.HexColor("#1A2B4C"))
    c.setFont(F_SANS_BOLD, 10.5)
    c.drawString(mx, my, "PROFESSIONAL WORK EXPERIENCE")
    my -= 5
    c.line(mx, my, mx + mw, my)
    my -= 14

    for exp in MARCUS_PROFILE["experience"]:
        c.setFont(F_SANS_BOLD, 9)
        c.setFillColor(colors.HexColor("#0F172A"))
        c.drawString(mx, my, exp["position"].upper())
        c.setFont(F_SANS_BOLD, 8)
        c.setFillColor(colors.HexColor("#0284C7"))
        date_str = exp["duration"]
        c.drawRightString(mx + mw, my, date_str)
        my -= 10

        c.setFont(F_SANS_ITALIC, 8)
        c.setFillColor(colors.HexColor("#475569"))
        c.drawString(mx, my, f"{exp['company']} — {exp['location']}")
        my -= 10

        for bullet in exp["bullets"]:
            bullet_txt = f"•  {bullet}"
            my = draw_wrapped(c, bullet_txt, mx + 5, my, mw - 5, F_SANS, 8, colors.HexColor("#334155"), 10.5)
        my -= 6

    c.save()
    return output_pdf


# ======================================================================
# RESUME DESIGN B (DOCX): Forest Green & Gold Top-Banner Layout
# Target Role: Hotel Events Service Coordinator
# ======================================================================
def build_marcus_resume2_docx(output_docx):
    doc = docx.Document()
    for sec in doc.sections:
        sec.top_margin = Inches(0.65)
        sec.bottom_margin = Inches(0.65)
        sec.left_margin = Inches(0.75)
        sec.right_margin = Inches(0.75)

    def add_sec_head(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(title.upper())
        run.font.name = 'Georgia'
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1B, 0x4D, 0x3E)
        pBorder = parse_xml(r'<w:pBdr {}><w:bottom w:val="single" w:sz="8" w:space="2" w:color="1B4D3E"/></w:pBdr>'.format(nsdecls('w')))
        p._p.get_or_add_pPr().append(pBorder)

    # Header
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_after = Pt(2)
    p_name.paragraph_format.space_before = Pt(0)
    run_n = p_name.add_run("MARCUS ELIJAH NAVARRO")
    run_n.font.name = 'Georgia'
    run_n.font.size = Pt(18)
    run_n.font.bold = True
    run_n.font.color.rgb = RGBColor(0x1B, 0x4D, 0x3E)

    p_pos = doc.add_paragraph()
    p_pos.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_pos.paragraph_format.space_after = Pt(4)
    run_p = p_pos.add_run("HOTEL EVENTS SERVICE COORDINATOR")
    run_p.font.name = 'Calibri'
    run_p.font.size = Pt(11)
    run_p.font.bold = True
    run_p.font.color.rgb = RGBColor(0xB8, 0x86, 0x0B)

    p_con = doc.add_paragraph()
    p_con.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_con.paragraph_format.space_after = Pt(10)
    run_c = p_con.add_run(f"{MARCUS_PROFILE['address']} | {MARCUS_PROFILE['phone']} | {MARCUS_PROFILE['email']}")
    run_c.font.name = 'Calibri'
    run_c.font.size = Pt(8.5)
    run_c.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    add_sec_head("Professional Event Coordination Summary")
    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.space_after = Pt(6)
    p_sum.paragraph_format.line_spacing = 1.15
    run_s = p_sum.add_run(
        "Detail-oriented Hotel Events Service Coordinator with comprehensive expertise in orchestrating corporate "
        "conferences, state banquets, and luxury weddings. Accomplished in bridging client expectations with hotel operational "
        "departments, verifying Banquet Event Order (BEO) parameters, supervising service timing, and managing floor setup "
        "crews. Proven ability to oversee simultaneous function room operations while preserving high standards of hospitality."
    )
    run_s.font.name = 'Calibri'
    run_s.font.size = Pt(9.5)

    add_sec_head("Areas of Event Coordination Expertise")
    skills_events = [
        ["Banquet Event Order (BEO) Analysis & Execution", "Function Room Staging & Table Floor Plans"],
        ["Client Coordination & On-Site Requirement Tracking", "Cross-Departmental Kitchen & Service Alignment"],
        ["VIP Protocol, Guest Arrival & Seating Management", "Catering Service Flow & Course Rollout Timing"],
        ["Staff Shift Briefings & Banquet Crew Leadership", "Post-Event Billing, Settlement & Feedback Collection"]
    ]
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r in skills_events:
        row = table.add_row()
        for idx, item in enumerate(r):
            cell = row.cells[idx]
            cell.width = Inches(3.4)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            run = p.add_run(f"•  {item}")
            run.font.name = 'Calibri'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    add_sec_head("Professional Hospitality Experience")
    for exp in MARCUS_PROFILE["experience"]:
        p_role = doc.add_paragraph()
        p_role.paragraph_format.space_before = Pt(6)
        p_role.paragraph_format.space_after = Pt(1)
        r_title = p_role.add_run(exp["position"])
        r_title.bold = True
        r_title.font.name = 'Georgia'
        r_title.font.size = Pt(10)
        r_title.font.color.rgb = RGBColor(0x1B, 0x4D, 0x3E)

        r_sep = p_role.add_run(" | ")
        r_sep.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

        r_comp = p_role.add_run(f"{exp['company']} — {exp['location']}")
        r_comp.font.name = 'Calibri'
        r_comp.font.size = Pt(9.5)
        r_comp.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

        r_dur = p_role.add_run(f"  ({exp['duration']})")
        r_dur.font.name = 'Calibri'
        r_dur.font.size = Pt(9)
        r_dur.font.bold = True
        r_dur.font.color.rgb = RGBColor(0xB8, 0x86, 0x0B)

        for b in exp["bullets"]:
            p_b = doc.add_paragraph(style='List Bullet')
            p_b.paragraph_format.space_after = Pt(2)
            p_b.paragraph_format.space_before = Pt(0)
            p_b.paragraph_format.line_spacing = 1.1
            run_b = p_b.add_run(b)
            run_b.font.name = 'Calibri'
            run_b.font.size = Pt(9)

    add_sec_head("Education & Professional Credentials")
    p_edu = doc.add_paragraph()
    p_edu.paragraph_format.space_after = Pt(2)
    r_ed = p_edu.add_run(f"{MARCUS_PROFILE['education']['degree']} (Graduated: {MARCUS_PROFILE['education']['year']})")
    r_ed.bold = True
    r_ed.font.name = 'Calibri'
    r_ed.font.size = Pt(9)
    p_inst = doc.add_paragraph()
    p_inst.paragraph_format.space_after = Pt(4)
    r_in = p_inst.add_run(MARCUS_PROFILE['education']['institution'])
    r_in.font.name = 'Calibri'
    r_in.font.size = Pt(8.5)
    r_in.font.italic = True

    for c_item in MARCUS_PROFILE["certifications"]:
        p_c = doc.add_paragraph(style='List Bullet')
        p_c.paragraph_format.space_after = Pt(1)
        r_c = p_c.add_run(f"{c_item['title']} — {c_item['issuer']} ({c_item['date']})")
        r_c.font.name = 'Calibri'
        r_c.font.size = Pt(8.5)

    doc.save(safe_path(output_docx))
    return output_docx


# ======================================================================
# RESUME DESIGN C (PNG): Burgundy Top Header Banner & 2-Col Layout
# Target Role: Banquet Team Leader
# ======================================================================
def build_marcus_resume3_png(output_png):
    temp_pdf = output_png.replace(".png", "_temp.pdf")
    c = canvas.Canvas(safe_path(temp_pdf), pagesize=letter)
    w, h = letter

    c.setFillColor(colors.HexColor("#722F37"))
    c.rect(0, h - 90, w, 90, stroke=0, fill=1)

    c.setFillColor(colors.HexColor("#D4AF37"))
    c.rect(0, h - 94, w, 4, stroke=0, fill=1)

    c.setFillColor(colors.white)
    c.setFont(F_SANS_BOLD, 17)
    c.drawString(30, h - 45, "MARCUS ELIJAH NAVARRO")
    c.setFillColor(colors.HexColor("#FDE68A"))
    c.setFont(F_SANS_BOLD, 9.5)
    c.drawString(30, h - 65, "BANQUET TEAM LEADER — STAFF LEADERSHIP & FUNCTION EXECUTION")

    c.setFont(F_SANS, 7.5)
    c.setFillColor(colors.HexColor("#F1F5F9"))
    c.drawRightString(w - 30, h - 42, f"Phone: {MARCUS_PROFILE['phone']}")
    c.drawRightString(w - 30, h - 55, f"Email: {MARCUS_PROFILE['email']}")
    c.drawRightString(w - 30, h - 68, f"{MARCUS_PROFILE['address']}")

    left_w = 340
    right_x = 30 + left_w + 20
    right_w = w - right_x - 30
    curr_ly = h - 115
    curr_ry = h - 115

    c.setFillColor(colors.HexColor("#722F37"))
    c.setFont(F_SANS_BOLD, 10)
    c.drawString(30, curr_ly, "TEAM LEADERSHIP PROFILE")
    curr_ly -= 4
    c.setStrokeColor(colors.HexColor("#722F37"))
    c.setLineWidth(1)
    c.line(30, curr_ly, 30 + left_w, curr_ly)
    curr_ly -= 12

    sum_text = (
        "Action-oriented Banquet Team Leader with proven ability to motivate, brief, and coordinate banquet service "
        "crews during high-capacity banquets and galas. Adept at on-the-floor coaching, sequence of service timing, table "
        "turnovers, and breakage mitigation. Known for maintaining composure during peak rushes and upholding Forbes standards."
    )
    curr_ly = draw_wrapped(c, sum_text, 30, curr_ly, left_w, F_SANS, 8, colors.HexColor("#334155"), 11)

    curr_ly -= 10
    c.setFillColor(colors.HexColor("#722F37"))
    c.setFont(F_SANS_BOLD, 10)
    c.drawString(30, curr_ly, "HOSPITALITY LEADERSHIP EXPERIENCE")
    curr_ly -= 4
    c.line(30, curr_ly, 30 + left_w, curr_ly)
    curr_ly -= 12

    for exp in MARCUS_PROFILE["experience"]:
        c.setFont(F_SANS_BOLD, 8.5)
        c.setFillColor(colors.HexColor("#0F172A"))
        c.drawString(30, curr_ly, exp["position"].upper())
        curr_ly -= 10
        c.setFont(F_SANS_ITALIC, 7.5)
        c.setFillColor(colors.HexColor("#722F37"))
        c.drawString(30, curr_ly, f"{exp['company']} | {exp['duration']}")
        curr_ly -= 10

        for b in exp["bullets"]:
            b_txt = f"•  {b}"
            curr_ly = draw_wrapped(c, b_txt, 35, curr_ly, left_w - 5, F_SANS, 7.5, colors.HexColor("#334155"), 10)
        curr_ly -= 6

    c.setFillColor(colors.HexColor("#722F37"))
    c.setFont(F_SANS_BOLD, 10)
    c.drawString(right_x, curr_ry, "CORE SKILLS & METRICS")
    curr_ry -= 4
    c.setStrokeColor(colors.HexColor("#722F37"))
    c.line(right_x, curr_ry, right_x + right_w, curr_ry)
    curr_ry -= 12

    skills_team = [
        "Pre-Shift Staff Briefings", "Service Brigade Leadership", "BEO Floor Timing Control",
        "VIP Table Protocol", "Breakage Control & Par Levels", "HACCP Food Safety",
        "Beverage Requisition Audit", "Banquet Hall Turnovers"
    ]
    for sk in skills_team:
        c.setFillColor(colors.HexColor("#FDF2F8"))
        c.rect(right_x, curr_ry - 2, right_w, 15, stroke=0, fill=1)
        c.setFillColor(colors.HexColor("#722F37"))
        c.setFont(F_SANS_BOLD, 7.5)
        c.drawString(right_x + 6, curr_ry + 2, f"✔  {sk}")
        curr_ry -= 18

    curr_ry -= 8
    c.setFillColor(colors.HexColor("#722F37"))
    c.setFont(F_SANS_BOLD, 10)
    c.drawString(right_x, curr_ry, "EDUCATION")
    curr_ry -= 4
    c.line(right_x, curr_ry, right_x + right_w, curr_ry)
    curr_ry -= 12

    c.setFont(F_SANS_BOLD, 8)
    c.setFillColor(colors.HexColor("#1E293B"))
    c.drawString(right_x, curr_ry, "BS Hospitality Management")
    curr_ry -= 9
    c.setFont(F_SANS, 7)
    c.setFillColor(colors.HexColor("#64748B"))
    c.drawString(right_x, curr_ry, "PLM – Intramuros (2018)")
    curr_ry -= 18

    c.setFillColor(colors.HexColor("#722F37"))
    c.setFont(F_SANS_BOLD, 10)
    c.drawString(right_x, curr_ry, "CREDENTIALS")
    curr_ry -= 4
    c.line(right_x, curr_ry, right_x + right_w, curr_ry)
    curr_ry -= 12

    for ct in MARCUS_PROFILE["certifications"]:
        c.setFont(F_SANS_BOLD, 7.5)
        c.setFillColor(colors.HexColor("#0F172A"))
        c.drawString(right_x, curr_ry, ct["title"])
        curr_ry -= 9
        c.setFont(F_SANS, 7)
        c.setFillColor(colors.HexColor("#64748B"))
        c.drawString(right_x, curr_ry, f"{ct['issuer']} ({ct['date']})")
        curr_ry -= 12

    c.save()
    pdf_to_image(temp_pdf, output_png, scale=2, format="PNG")
    if os.path.exists(safe_path(temp_pdf)):
        os.remove(safe_path(temp_pdf))
    return output_png


# ======================================================================
# RESUME DESIGN D (JPG): Warm Bronze & Metric Cards Minimalist Grid
# Target Role: Catering Service Supervisor
# ======================================================================
def build_marcus_resume4_jpg(output_jpg):
    temp_pdf = output_jpg.replace(".jpg", "_temp.pdf")
    c = canvas.Canvas(safe_path(temp_pdf), pagesize=letter)
    w, h = letter

    c.setFillColor(colors.HexColor("#854D0E"))
    c.rect(0, h - 8, w, 8, stroke=0, fill=1)

    c.setFillColor(colors.HexColor("#1E293B"))
    c.setFont(F_SERIF_BOLD, 19)
    c.drawString(35, h - 45, "MARCUS ELIJAH NAVARRO")
    c.setFont(F_SERIF_ITALIC, 10.5)
    c.setFillColor(colors.HexColor("#854D0E"))
    c.drawString(35, h - 62, "CATERING SERVICE SUPERVISOR — CATERING & F&B EVENT OPERATIONS")

    c.setFont(F_SANS, 7.5)
    c.setFillColor(colors.HexColor("#64748B"))
    c.drawString(35, h - 76, f"{MARCUS_PROFILE['address']} | {MARCUS_PROFILE['phone']} | {MARCUS_PROFILE['email']}")

    card_w = (w - 70 - 20) / 3
    card_y = h - 128
    metrics = [
        ("600+ CAPACITY", "Managed Banquet Functions"),
        ("18% REDUCTION", "In Equipment & Glass Loss"),
        ("5+ YEARS", "Hospitality Catering Leadership")
    ]
    for idx, (m_val, m_lbl) in enumerate(metrics):
        cx = 35 + idx * (card_w + 10)
        c.setFillColor(colors.HexColor("#FEFCE8"))
        c.rect(cx, card_y, card_w, 40, stroke=1, fill=1)
        c.setStrokeColor(colors.HexColor("#FEF08A"))
        c.setFillColor(colors.HexColor("#854D0E"))
        c.setFont(F_SANS_BOLD, 10)
        c.drawCentredString(cx + card_w / 2, card_y + 24, m_val)
        c.setFillColor(colors.HexColor("#713F12"))
        c.setFont(F_SANS, 7.5)
        c.drawCentredString(cx + card_w / 2, card_y + 10, m_lbl)

    curr_y = card_y - 20
    c.setFont(F_SERIF_BOLD, 11)
    c.setFillColor(colors.HexColor("#854D0E"))
    c.drawString(35, curr_y, "Catering Management Profile")
    curr_y -= 4
    c.setStrokeColor(colors.HexColor("#CA8A04"))
    c.line(35, curr_y, w - 35, curr_y)
    curr_y -= 12

    cater_summary = (
        "Seasoned Catering Service Supervisor with substantial background leading high-volume culinary events, "
        "buffet logistics, and VIP banquet receptions. Expert in organizing buffet line staging, table appointments, "
        "equipment allocation, and temperature monitoring. Proven record in minimizing chinaware breakage, ensuring "
        "prompt replenishment, and maintaining seamless coordination between kitchen and banquet floors."
    )
    curr_y = draw_wrapped(c, cater_summary, 35, curr_y, w - 70, F_SERIF, 8.5, colors.HexColor("#334155"), 11.5)

    curr_y -= 10
    c.setFont(F_SERIF_BOLD, 11)
    c.setFillColor(colors.HexColor("#854D0E"))
    c.drawString(35, curr_y, "Catering & Banquet Experience")
    curr_y -= 4
    c.line(35, curr_y, w - 35, curr_y)
    curr_y -= 14

    for exp in MARCUS_PROFILE["experience"]:
        c.setFont(F_SERIF_BOLD, 9)
        c.setFillColor(colors.HexColor("#0F172A"))
        c.drawString(35, curr_y, exp["position"])
        c.setFont(F_SANS_BOLD, 8)
        c.setFillColor(colors.HexColor("#854D0E"))
        c.drawRightString(w - 35, curr_y, exp["duration"])
        curr_y -= 10

        c.setFont(F_SERIF_ITALIC, 8)
        c.setFillColor(colors.HexColor("#475569"))
        c.drawString(35, curr_y, f"{exp['company']} — {exp['location']}")
        curr_y -= 10

        for b in exp["bullets"]:
            b_txt = f"•  {b}"
            curr_y = draw_wrapped(c, b_txt, 42, curr_y, w - 77, F_SANS, 8, colors.HexColor("#334155"), 10.5)
        curr_y -= 6

    curr_y -= 6
    c.setFont(F_SERIF_BOLD, 10)
    c.setFillColor(colors.HexColor("#854D0E"))
    c.drawString(35, curr_y, "Academic Background & Certifications")
    curr_y -= 4
    c.line(35, curr_y, w - 35, curr_y)
    curr_y -= 12

    c.setFont(F_SANS_BOLD, 8)
    c.setFillColor(colors.HexColor("#1E293B"))
    c.drawString(35, curr_y, f"Degree: {MARCUS_PROFILE['education']['degree']} — {MARCUS_PROFILE['education']['institution']} ({MARCUS_PROFILE['education']['year']})")
    curr_y -= 11

    certs_line = " | ".join([f"{c_item['title']} ({c_item['issuer']})" for c_item in MARCUS_PROFILE["certifications"][:2]])
    c.setFont(F_SANS, 7.5)
    c.setFillColor(colors.HexColor("#475569"))
    c.drawString(35, curr_y, f"Certifications: {certs_line}")

    c.save()
    pdf_to_image(temp_pdf, output_jpg, scale=2, format="JPG", quality=92)
    if os.path.exists(safe_path(temp_pdf)):
        os.remove(safe_path(temp_pdf))
    return output_jpg


# ======================================================================
# RESUME DESIGN E (Blurred PNG): Slate Teal Split Header & Card Layout
# Target Role: Function Operations Officer
# ======================================================================
def build_marcus_resume5_blurred_png(output_png):
    temp_clean_png = output_png.replace(".png", "_clean.png")
    temp_pdf = output_png.replace(".png", "_temp.pdf")
    c = canvas.Canvas(safe_path(temp_pdf), pagesize=letter)
    w, h = letter

    c.setFillColor(colors.HexColor("#0F766E"))
    c.rect(0, h - 75, w, 75, stroke=0, fill=1)

    c.setFillColor(colors.white)
    c.setFont(F_SANS_BOLD, 16)
    c.drawString(30, h - 38, "MARCUS ELIJAH NAVARRO")
    c.setFillColor(colors.HexColor("#99F6E4"))
    c.setFont(F_SANS_BOLD, 9)
    c.drawString(30, h - 55, "FUNCTION OPERATIONS OFFICER — OPERATIONAL COORDINATION")

    c.setFont(F_SANS, 7)
    c.setFillColor(colors.white)
    c.drawRightString(w - 30, h - 35, MARCUS_PROFILE["phone"])
    c.drawRightString(w - 30, h - 47, MARCUS_PROFILE["email"])
    c.drawRightString(w - 30, h - 59, "Pasay City, Metro Manila")

    curr_y = h - 95

    c.setFillColor(colors.HexColor("#F0FDFA"))
    c.rect(30, curr_y - 45, w - 60, 45, stroke=1, fill=1)
    c.setStrokeColor(colors.HexColor("#CCFBF1"))
    c.setFont(F_SANS_BOLD, 8.5)
    c.setFillColor(colors.HexColor("#0F766E"))
    c.drawString(38, curr_y - 12, "FUNCTION ROOM OPERATIONS PROFILE:")
    func_sum = (
        "Rigorous Function Operations Officer specializing in room setups, turnover schedules, audiovisual alignment, "
        "and floor safety for luxury hotel conventions and private functions. Committed to rapid room transformations, "
        "equipment inventory integrity, and cross-departmental coordination."
    )
    draw_wrapped(c, func_sum, 38, curr_y - 23, w - 76, F_SANS, 7.5, colors.HexColor("#334155"), 9.5)

    curr_y -= 60

    c.setFont(F_SANS_BOLD, 9.5)
    c.setFillColor(colors.HexColor("#0F766E"))
    c.drawString(30, curr_y, "FUNCTION OPERATIONAL COMPETENCIES")
    curr_y -= 4
    c.setStrokeColor(colors.HexColor("#0F766E"))
    c.line(30, curr_y, w - 30, curr_y)
    curr_y -= 12

    sk_list = [
        "Room Turnover Logistics", "Audiovisual & Staging Check", "BEO Floor Specs Verification",
        "Crowd Flow & Safety Protocols", "Table Layout Drafting", "Equipment Par Level Counts",
        "Beverage Requisition Tracking", "VIP Function Setups", "Breakage Control Auditing"
    ]
    col_w = (w - 60) / 3
    for r in range(3):
        for col in range(3):
            idx = r * 3 + col
            c.setFont(F_SANS, 7.5)
            c.setFillColor(colors.HexColor("#1E293B"))
            c.drawString(30 + col * col_w, curr_y, f"▪  {sk_list[idx]}")
        curr_y -= 11

    curr_y -= 8
    c.setFont(F_SANS_BOLD, 9.5)
    c.setFillColor(colors.HexColor("#0F766E"))
    c.drawString(30, curr_y, "CHRONOLOGICAL HOSPITALITY WORK EXPERIENCE")
    curr_y -= 4
    c.line(30, curr_y, w - 30, curr_y)
    curr_y -= 14

    for exp in MARCUS_PROFILE["experience"]:
        c.setFillColor(colors.HexColor("#0F766E"))
        c.setFont(F_SANS_BOLD, 8.5)
        c.drawString(30, curr_y, exp["position"].upper())
        c.setFont(F_SANS_BOLD, 7.5)
        c.setFillColor(colors.HexColor("#0D9488"))
        c.drawRightString(w - 30, curr_y, exp["duration"])
        curr_y -= 9

        c.setFont(F_SANS_ITALIC, 7.5)
        c.setFillColor(colors.HexColor("#64748B"))
        c.drawString(30, curr_y, f"{exp['company']} — {exp['location']}")
        curr_y -= 9

        for b in exp["bullets"]:
            b_txt = f"•  {b}"
            curr_y = draw_wrapped(c, b_txt, 36, curr_y, w - 66, F_SANS, 7.5, colors.HexColor("#334155"), 9.5)
        curr_y -= 5

    curr_y -= 6
    c.setFont(F_SANS_BOLD, 8.5)
    c.setFillColor(colors.HexColor("#0F766E"))
    c.drawString(30, curr_y, "ACADEMIC & CERTIFICATION CREDENTIALS")
    curr_y -= 3
    c.line(30, curr_y, w - 30, curr_y)
    curr_y -= 10

    c.setFont(F_SANS, 7.5)
    c.setFillColor(colors.HexColor("#1E293B"))
    c.drawString(30, curr_y, f"Education: {MARCUS_PROFILE['education']['degree']} | {MARCUS_PROFILE['education']['institution']} ({MARCUS_PROFILE['education']['year']})")
    curr_y -= 10
    c_str = ", ".join([c_item["title"] for c_item in MARCUS_PROFILE["certifications"]])
    c.drawString(30, curr_y, f"Certifications: {c_str}")

    c.save()
    pdf_to_image(temp_pdf, temp_clean_png, scale=2, format="PNG")
    apply_uniform_blur(temp_clean_png, output_png, blur_radius=1.8, format="PNG")
    if os.path.exists(safe_path(temp_pdf)):
        os.remove(safe_path(temp_pdf))
    if os.path.exists(safe_path(temp_clean_png)):
        os.remove(safe_path(temp_clean_png))
    return output_png


# ======================================================================
# RESUME DESIGN F (Blurred JPG): Royal Purple Classic Framed Layout
# Target Role: Hotel Events Operations Coordinator
# ======================================================================
def build_marcus_resume6_blurred_jpg(output_jpg):
    temp_clean_jpg = output_jpg.replace(".jpg", "_clean.jpg")
    temp_pdf = output_jpg.replace(".jpg", "_temp.pdf")
    c = canvas.Canvas(safe_path(temp_pdf), pagesize=letter)
    w, h = letter

    c.setStrokeColor(colors.HexColor("#4A0E4E"))
    c.setLineWidth(2)
    c.rect(20, 20, w - 40, h - 40)
    c.setStrokeColor(colors.HexColor("#9C27B0"))
    c.setLineWidth(0.8)
    c.rect(24, 24, w - 48, h - 48)

    c.setFillColor(colors.HexColor("#4A0E4E"))
    c.setFont(F_SERIF_BOLD, 17)
    c.drawCentredString(w / 2, h - 50, "MARCUS ELIJAH NAVARRO")
    c.setFont(F_SERIF_ITALIC, 9.5)
    c.setFillColor(colors.HexColor("#7B1FA2"))
    c.drawCentredString(w / 2, h - 65, "HOTEL EVENTS OPERATIONS COORDINATOR")
    c.setFont(F_SANS, 7.5)
    c.setFillColor(colors.HexColor("#555555"))
    c.drawCentredString(w / 2, h - 77, f"{MARCUS_PROFILE['address']} • {MARCUS_PROFILE['phone']} • {MARCUS_PROFILE['email']}")

    curr_y = h - 96

    def sec_head(title):
        nonlocal curr_y
        c.setFillColor(colors.HexColor("#4A0E4E"))
        c.setFont(F_SERIF_BOLD, 9.5)
        c.drawString(35, curr_y, title.upper())
        curr_y -= 3
        c.setStrokeColor(colors.HexColor("#BA68C8"))
        c.setLineWidth(0.8)
        c.line(35, curr_y, w - 35, curr_y)
        curr_y -= 11

    sec_head("Hotel Events Operations Summary")
    ev_sum = (
        "Accomplished Hotel Events Operations Coordinator dedicated to the high-standard execution of banquets, "
        "conventions, and luxury celebrations. Proficient in guest experience monitoring, banquet floor sequence "
        "enforcement, staff supervision, and inter-departmental event alignment. Demonstrates a strong commitment to Forbes "
        "standards of hospitality and seamless guest satisfaction."
    )
    curr_y = draw_wrapped(c, ev_sum, 35, curr_y, w - 70, F_SERIF, 8, colors.HexColor("#333333"), 10.5)

    curr_y -= 8
    sec_head("Core Operational Competencies")
    skills = [
        "• Banquet Event Order (BEO) Floor Alignment", "• Function Setup & Room Turnover Flow",
        "• Banquet Team Shift Briefings (up to 25)", "• VIP Protocol & Guest Experience Standards",
        "• Plated Service Timing & Buffet Management", "• Chinaware & Equipment Inventory Control",
        "• HACCP Food Safety & Sanitation Rules", "• Event Settlement & Beverage Requisition"
    ]
    half_w = (w - 70) / 2
    for i in range(0, len(skills), 2):
        c.setFont(F_SANS, 7.5)
        c.setFillColor(colors.HexColor("#222222"))
        c.drawString(35, curr_y, skills[i])
        c.drawString(35 + half_w, curr_y, skills[i + 1])
        curr_y -= 10

    curr_y -= 8
    sec_head("Chronological Professional Experience")
    for exp in MARCUS_PROFILE["experience"]:
        c.setFont(F_SERIF_BOLD, 8.5)
        c.setFillColor(colors.HexColor("#4A0E4E"))
        c.drawString(35, curr_y, exp["position"])
        c.setFont(F_SANS_BOLD, 7.5)
        c.setFillColor(colors.HexColor("#7B1FA2"))
        c.drawRightString(w - 35, curr_y, exp["duration"])
        curr_y -= 9

        c.setFont(F_SERIF_ITALIC, 7.5)
        c.setFillColor(colors.HexColor("#555555"))
        c.drawString(35, curr_y, f"{exp['company']} — {exp['location']}")
        curr_y -= 9

        for b in exp["bullets"]:
            b_txt = f"•  {b}"
            curr_y = draw_wrapped(c, b_txt, 42, curr_y, w - 77, F_SANS, 7.5, colors.HexColor("#333333"), 9.5)
        curr_y -= 4

    curr_y -= 6
    sec_head("Education & Professional Accreditations")
    c.setFont(F_SERIF_BOLD, 8)
    c.setFillColor(colors.HexColor("#222222"))
    c.drawString(35, curr_y, f"{MARCUS_PROFILE['education']['degree']} — {MARCUS_PROFILE['education']['institution']} ({MARCUS_PROFILE['education']['year']})")
    curr_y -= 10

    for ct in MARCUS_PROFILE["certifications"]:
        c.setFont(F_SANS, 7)
        c.setFillColor(colors.HexColor("#444444"))
        c.drawString(42, curr_y, f"•  {ct['title']} — {ct['issuer']} ({ct['date']})")
        curr_y -= 9

    c.save()
    pdf_to_image(temp_pdf, temp_clean_jpg, scale=2, format="JPG", quality=92)
    apply_uniform_blur(temp_clean_jpg, output_jpg, blur_radius=1.8, format="JPG", quality=85)
    if os.path.exists(safe_path(temp_pdf)):
        os.remove(safe_path(temp_pdf))
    if os.path.exists(safe_path(temp_clean_jpg)):
        os.remove(safe_path(temp_clean_jpg))
    return output_jpg


# ======================================================================
# SUPPORTING DOCUMENTS FOR MARCUS ELIJAH NAVARRO (6 Documents)
# ======================================================================

def build_marcus_doc1_coe_current(output_pdf):
    c = canvas.Canvas(safe_path(output_pdf), pagesize=letter)
    w, h = letter

    c.setFillColor(colors.HexColor("#1A2B4C"))
    c.rect(0, h - 85, w, 85, stroke=0, fill=1)
    c.setFillColor(colors.HexColor("#D4AF37"))
    c.setFont(F_SERIF_BOLD, 18)
    c.drawString(35, h - 42, "GRAND CREST HOTEL & SUITES")
    c.setFont(F_SANS, 8)
    c.setFillColor(colors.HexColor("#E2E8F0"))
    c.drawString(35, h - 58, "Luxury Accommodation & World-Class Convention Centre")
    c.drawRightString(w - 35, h - 42, "Seaside Boulevard, Pasay City 1308")
    c.drawRightString(w - 35, h - 54, "Metro Manila, Philippines")
    c.drawRightString(w - 35, h - 66, "Tel: +63 2 8888 7000 | hr@grandcrestmanila.com")

    curr_y = h - 120
    c.setFillColor(colors.HexColor("#1E293B"))
    c.setFont(F_SANS_BOLD, 8.5)
    c.drawString(35, curr_y, "DATE OF ISSUANCE: October 2, 2026")
    curr_y -= 12
    c.drawString(35, curr_y, "REF NO: GCHS-HRD-COE-2026-0842")
    curr_y -= 25

    c.setFont(F_SERIF_BOLD, 13)
    c.setFillColor(colors.HexColor("#1A2B4C"))
    c.drawCentredString(w / 2, curr_y, "CERTIFICATE OF EMPLOYMENT")
    curr_y -= 6
    c.setStrokeColor(colors.HexColor("#D4AF37"))
    c.setLineWidth(1)
    c.line(w / 2 - 120, curr_y, w / 2 + 120, curr_y)
    curr_y -= 25

    body1 = (
        "TO WHOM IT MAY CONCERN:\n\n"
        "This is to certify that MR. MARCUS ELIJAH NAVARRO is a bona fide employee of Grand Crest Hotel & Suites, "
        "holding the position of Banquet Operations Supervisor from June 2022 up to the present date.\n\n"
        "In his supervisory capacity, Mr. Navarro is responsible for overseeing ballroom and multi-function banquet operations, "
        "directing banquet captain briefings, coordinating Banquet Event Orders (BEOs) with culinary and banquet sales brigades, "
        "and upholding HACCP food safety standards across high-profile conventions and wedding banquets.\n\n"
        "Mr. Navarro has demonstrated exceptional leadership, operational discipline, and commitment to luxury guest service. "
        "He maintains an exemplary employment record with zero disciplinary infractions.\n\n"
        "This certification is issued upon the request of Mr. Navarro for employment verification and professional credential evaluation purposes."
    )
    for p in body1.split('\n\n'):
        curr_y = draw_wrapped(c, p, 35, curr_y, w - 70, F_SERIF, 9.5, colors.HexColor("#1E293B"), 13.5)
        curr_y -= 10

    curr_y -= 20
    c.setFont(F_SERIF_BOLD, 9.5)
    c.drawString(35, curr_y, "ELEANOR MONTEFALCO, MBA, CHRM")
    curr_y -= 11
    c.setFont(F_SANS, 8.5)
    c.setFillColor(colors.HexColor("#64748B"))
    c.drawString(35, curr_y, "Director of Human Resources & Talent Development")
    curr_y -= 10
    c.drawString(35, curr_y, "Grand Crest Hotel & Suites Manila")

    c.save()
    return output_pdf


def build_marcus_doc2_coe_previous(output_pdf):
    c = canvas.Canvas(safe_path(output_pdf), pagesize=letter)
    w, h = letter

    c.setFillColor(colors.HexColor("#0C4A6E"))
    c.rect(0, h - 80, w, 80, stroke=0, fill=1)
    c.setFillColor(colors.HexColor("#38BDF8"))
    c.setFont(F_SERIF_BOLD, 17)
    c.drawString(35, h - 40, "HARBORVIEW HOTEL & CONVENTION CENTER")
    c.setFont(F_SANS, 8)
    c.setFillColor(colors.white)
    c.drawString(35, h - 55, "Roxas Boulevard, Manila Bay, Metro Manila, Philippines")
    c.drawRightString(w - 35, h - 45, "Human Capital Division")
    c.drawRightString(w - 35, h - 57, "Email: careers@harborviewhotel.ph")

    curr_y = h - 115
    c.setFillColor(colors.HexColor("#1E293B"))
    c.setFont(F_SANS_BOLD, 8.5)
    c.drawString(35, curr_y, "DATE: May 30, 2022")
    curr_y -= 12
    c.drawString(35, curr_y, "REFERENCE ID: HVC-EMP-2022-319")
    curr_y -= 25

    c.setFont(F_SERIF_BOLD, 12.5)
    c.setFillColor(colors.HexColor("#0C4A6E"))
    c.drawCentredString(w / 2, curr_y, "CERTIFICATE OF EMPLOYMENT AND CLEARANCE")
    curr_y -= 6
    c.setStrokeColor(colors.HexColor("#38BDF8"))
    c.setLineWidth(1)
    c.line(w / 2 - 140, curr_y, w / 2 + 140, curr_y)
    curr_y -= 25

    body = (
        "TO WHOM IT MAY CONCERN:\n\n"
        "This is to formally certify that MARCUS ELIJAH NAVARRO was employed by Harborview Hotel & Convention Center "
        "from January 15, 2021 to May 20, 2022.\n\n"
        "During his tenure, Mr. Navarro held the position of Banquet Team Leader / Captain in the Food and Beverage "
        "Department. His duties included supervising banquet attendants during diplomatic functions, inspecting function room "
        "setups, executing Banquet Event Orders, and conducting equipment par level audits.\n\n"
        "Mr. Navarro has completed his formal exit clearance process and has been relieved of all financial and property "
        "accountabilities with Harborview Hotel & Convention Center.\n\n"
        "This certification is issued upon his request for whatever legal purpose it may serve."
    )
    for p in body.split('\n\n'):
        curr_y = draw_wrapped(c, p, 35, curr_y, w - 70, F_SERIF, 9.5, colors.HexColor("#1E293B"), 13.5)
        curr_y -= 10

    curr_y -= 25
    c.setFont(F_SERIF_BOLD, 9.5)
    c.drawString(35, curr_y, "VICTORIA R. CASTILLO, FHRM")
    curr_y -= 11
    c.setFont(F_SANS, 8.5)
    c.setFillColor(colors.HexColor("#64748B"))
    c.drawString(35, curr_y, "Vice President for Human Resources")
    curr_y -= 10
    c.drawString(35, curr_y, "Harborview Hotel & Convention Center")

    c.save()
    return output_pdf


def build_marcus_doc3_diploma(output_pdf):
    c = canvas.Canvas(safe_path(output_pdf), pagesize=landscape(letter))
    w, h = landscape(letter)

    c.setStrokeColor(colors.HexColor("#1E3A8A"))
    c.setLineWidth(4)
    c.rect(24, 24, w - 48, h - 48)
    c.setStrokeColor(colors.HexColor("#D4AF37"))
    c.setLineWidth(1.5)
    c.rect(30, 30, w - 60, h - 60)

    c.setFillColor(colors.HexColor("#1E3A8A"))
    c.setFont(F_SERIF_BOLD, 12)
    c.drawCentredString(w / 2, h - 70, "REPUBLIKA NG PILIPINAS")
    c.setFont(F_SERIF_BOLD, 18)
    c.drawCentredString(w / 2, h - 95, "PAMANTASAN NG LUNGSOD NG MAYNILA")
    c.setFont(F_SERIF_ITALIC, 10)
    c.setFillColor(colors.HexColor("#475569"))
    c.drawCentredString(w / 2, h - 110, "Intramuros, Maynila, Pilipinas")

    c.setFont(F_SERIF, 10.5)
    c.setFillColor(colors.HexColor("#1E293B"))
    c.drawCentredString(w / 2, h - 150, "Ang Lupon ng mga Rehente at ang Dalubhasaan ng Pamamahala sa Mabuting Pakikitungo")
    c.drawCentredString(w / 2, h - 168, "ay nagpapatunay na si")

    c.setFont(F_SERIF_BOLD, 22)
    c.setFillColor(colors.HexColor("#1E3A8A"))
    c.drawCentredString(w / 2, h - 210, "MARCUS ELIJAH NAVARRO")
    c.setStrokeColor(colors.HexColor("#D4AF37"))
    c.line(w / 2 - 180, h - 216, w / 2 + 180, h - 216)

    c.setFont(F_SERIF, 10.5)
    c.setFillColor(colors.HexColor("#1E293B"))
    c.drawCentredString(w / 2, h - 245, "ay kasiya-siyang nakatupad sa lahat ng mga kinakailangan para sa titulong")

    c.setFont(F_SERIF_BOLD, 15)
    c.setFillColor(colors.HexColor("#B45309"))
    c.drawCentredString(w / 2, h - 275, "BACHELOR OF SCIENCE IN HOSPITALITY MANAGEMENT")

    c.setFont(F_SERIF, 10)
    c.setFillColor(colors.HexColor("#334155"))
    c.drawCentredString(w / 2, h - 310, "Ipinagkaloob sa Lungsod ng Maynila, Pilipinas ngayong ika-18 ng Abril, 2018.")

    c.setFont(F_SERIF_BOLD, 9.5)
    c.setFillColor(colors.HexColor("#0F172A"))
    c.drawCentredString(w / 4, 90, "DR. MA. CECILIA S. ALVAREZ")
    c.setFont(F_SANS, 7.5)
    c.setFillColor(colors.HexColor("#64748B"))
    c.drawCentredString(w / 4, 76, "Dekana, Dalubhasaan ng Mabuting Pakikitungo")

    c.setFont(F_SERIF_BOLD, 9.5)
    c.setFillColor(colors.HexColor("#0F172A"))
    c.drawCentredString((3 * w) / 4, 90, "ATTY. EMMANUEL T. LORENZO")
    c.setFont(F_SANS, 7.5)
    c.setFillColor(colors.HexColor("#64748B"))
    c.drawCentredString((3 * w) / 4, 76, "Pangulo ng Pamantasan")

    c.save()
    return output_pdf


def build_marcus_doc4_cert_foodsafety(output_pdf):
    c = canvas.Canvas(safe_path(output_pdf), pagesize=landscape(letter))
    w, h = landscape(letter)

    c.setStrokeColor(colors.HexColor("#065F46"))
    c.setLineWidth(3.5)
    c.rect(26, 26, w - 52, h - 52)
    c.setStrokeColor(colors.HexColor("#10B981"))
    c.setLineWidth(1)
    c.rect(32, 32, w - 64, h - 64)

    c.setFillColor(colors.HexColor("#065F46"))
    c.setFont(F_SERIF_BOLD, 15)
    c.drawCentredString(w / 2, h - 70, "HOSPITALITY TRAINING COUNCIL OF THE PHILIPPINES")
    c.setFont(F_SANS, 8.5)
    c.setFillColor(colors.HexColor("#475569"))
    c.drawCentredString(w / 2, h - 85, "National Food Safety & Hygiene Accreditation Center — Pasig City, Philippines")

    c.setFont(F_SERIF_BOLD, 17)
    c.setFillColor(colors.HexColor("#065F46"))
    c.drawCentredString(w / 2, h - 130, "CERTIFICATE OF COMPETENCY")
    c.setFont(F_SERIF_ITALIC, 10.5)
    c.setFillColor(colors.HexColor("#334155"))
    c.drawCentredString(w / 2, h - 160, "This is to certify that")

    c.setFont(F_SERIF_BOLD, 20)
    c.setFillColor(colors.HexColor("#0F172A"))
    c.drawCentredString(w / 2, h - 195, "MARCUS ELIJAH NAVARRO")
    c.setStrokeColor(colors.HexColor("#065F46"))
    c.line(w / 2 - 160, h - 200, w / 2 + 160, h - 200)

    c.setFont(F_SERIF, 10)
    c.drawCentredString(w / 2, h - 225, "has satisfactorily completed the rigorous professional qualification requirements in")

    c.setFont(F_SERIF_BOLD, 13.5)
    c.setFillColor(colors.HexColor("#065F46"))
    c.drawCentredString(w / 2, h - 250, "FOOD SAFETY & SANITATION STANDARDS (HACCP LEVEL 2)")

    c.setFont(F_SANS, 8.5)
    c.setFillColor(colors.HexColor("#334155"))
    c.drawCentredString(w / 2, h - 280, "Covering Hazard Analysis Critical Control Point (HACCP) Verification, Food Temperature Controls,")
    c.drawCentredString(w / 2, h - 295, "Buffet Line Hygiene Protocols, and Banquet Sanitations (32 Instructional Hours).")

    c.drawCentredString(w / 2, h - 330, "Date Awarded: August 14, 2022  |  Certificate No: HTCP-HACCP2-2022-814")

    c.setFont(F_SERIF_BOLD, 9)
    c.setFillColor(colors.HexColor("#0F172A"))
    c.drawCentredString(w / 4, 85, "DR. ROLANDO S. GATMAITAN")
    c.setFont(F_SANS, 7.5)
    c.setFillColor(colors.HexColor("#64748B"))
    c.drawCentredString(w / 4, 72, "Lead HACCP Assessor & Food Hygiene Director")

    c.setFont(F_SERIF_BOLD, 9)
    c.setFillColor(colors.HexColor("#0F172A"))
    c.drawCentredString((3 * w) / 4, 85, "MARGARITA B. CORPUZ, CHA")
    c.setFont(F_SANS, 7.5)
    c.setFillColor(colors.HexColor("#64748B"))
    c.drawCentredString((3 * w) / 4, 72, "Executive Director, HTCP")

    c.save()
    return output_pdf


def build_marcus_doc5_cert_service(output_pdf):
    c = canvas.Canvas(safe_path(output_pdf), pagesize=landscape(letter))
    w, h = landscape(letter)

    c.setStrokeColor(colors.HexColor("#4F46E5"))
    c.setLineWidth(3)
    c.rect(28, 28, w - 56, h - 56)
    c.setStrokeColor(colors.HexColor("#A5B4FC"))
    c.setLineWidth(1)
    c.rect(34, 34, w - 68, h - 68)

    c.setFillColor(colors.HexColor("#4F46E5"))
    c.setFont(F_SERIF_BOLD, 15)
    c.drawCentredString(w / 2, h - 70, "PHILIPPINE HOSPITALITY DEVELOPMENT CENTER")
    c.setFont(F_SANS, 8.5)
    c.setFillColor(colors.HexColor("#475569"))
    c.drawCentredString(w / 2, h - 85, "Center for Banquet & Culinary Excellence — Quezon City, Philippines")

    c.setFont(F_SERIF_BOLD, 16)
    c.setFillColor(colors.HexColor("#4F46E5"))
    c.drawCentredString(w / 2, h - 130, "CERTIFICATE OF PARTICIPATION")

    c.setFont(F_SERIF_ITALIC, 10.5)
    c.setFillColor(colors.HexColor("#334155"))
    c.drawCentredString(w / 2, h - 160, "This certificate is proudly awarded to")

    c.setFont(F_SERIF_BOLD, 20)
    c.setFillColor(colors.HexColor("#1E1B4B"))
    c.drawCentredString(w / 2, h - 195, "MARCUS ELIJAH NAVARRO")
    c.setStrokeColor(colors.HexColor("#4F46E5"))
    c.line(w / 2 - 160, h - 200, w / 2 + 160, h - 200)

    c.setFont(F_SERIF, 10)
    c.drawCentredString(w / 2, h - 225, "for active participation and completion of the technical skills workshop in")

    c.setFont(F_SERIF_BOLD, 13.5)
    c.setFillColor(colors.HexColor("#4F46E5"))
    c.drawCentredString(w / 2, h - 250, "BASIC TABLE SETTING AND NAPKIN FOLDING WORKSHOP")

    c.setFont(F_SANS, 8.5)
    c.setFillColor(colors.HexColor("#334155"))
    c.drawCentredString(w / 2, h - 280, "Conducted on March 18, 2020 at the PHDC Training Suites, Quezon City.")
    c.drawCentredString(w / 2, h - 305, "Accreditation Code: PHDC-WS-2020-0931")

    c.setFont(F_SERIF_BOLD, 9)
    c.setFillColor(colors.HexColor("#0F172A"))
    c.drawCentredString(w / 2, 85, "FERNANDO D. MENDOZA, CHE")
    c.setFont(F_SANS, 7.5)
    c.setFillColor(colors.HexColor("#64748B"))
    c.drawCentredString(w / 2, 72, "Program Director, Philippine Hospitality Development Center")

    c.save()
    return output_pdf


def build_marcus_doc6_incomplete_cert(output_pdf):
    temp_clean_pdf = output_pdf.replace(".pdf", "_clean.pdf")
    c = canvas.Canvas(safe_path(temp_clean_pdf), pagesize=landscape(letter))
    w, h = landscape(letter)

    c.setStrokeColor(colors.HexColor("#854D0E"))
    c.setLineWidth(3)
    c.rect(26, 26, w - 52, h - 52)
    c.setStrokeColor(colors.HexColor("#EAB308"))
    c.setLineWidth(1)
    c.rect(32, 32, w - 64, h - 64)

    c.setFillColor(colors.HexColor("#854D0E"))
    c.setFont(F_SERIF_BOLD, 16)
    c.drawCentredString(w / 2, h - 70, "HOSPITALITY OPERATIONS INSTITUTE / FHCB")
    c.setFont(F_SANS, 8)
    c.setFillColor(colors.HexColor("#475569"))
    c.drawCentredString(w / 2, h - 85, "Federation of Hospitality Credentialing Boards")

    c.setFont(F_SERIF_BOLD, 17)
    c.setFillColor(colors.HexColor("#854D0E"))
    c.drawCentredString(w / 2, h - 130, "PROFESSIONAL SUPERVISORY DESIGNATION")

    c.setFont(F_SERIF_ITALIC, 10.5)
    c.setFillColor(colors.HexColor("#334155"))
    c.drawCentredString(w / 2, h - 160, "The Board of Directors hereby confers upon")

    c.setFont(F_SERIF_BOLD, 20)
    c.setFillColor(colors.HexColor("#1E293B"))
    c.drawCentredString(w / 2, h - 195, "MARCUS ELIJAH NAVARRO")

    c.setFont(F_SERIF, 10)
    c.drawCentredString(w / 2, h - 225, "the professional credential and title of")

    c.setFont(F_SERIF_BOLD, 15)
    c.setFillColor(colors.HexColor("#854D0E"))
    c.drawCentredString(w / 2, h - 250, "CERTIFIED HOSPITALITY SUPERVISOR (CHS)")

    c.setFont(F_SANS, 8.5)
    c.setFillColor(colors.HexColor("#475569"))
    c.drawCentredString(w / 2, h - 280, "Having successfully satisfied the experience, ethical benchmarks, and examination requirements.")
    c.drawCentredString(w / 2, h - 305, "Registration No: CHS-PH-2021-4198  |  Date of Issue: November 12, 2021")

    c.setFont(F_SERIF_BOLD, 9)
    c.setFillColor(colors.HexColor("#0F172A"))
    c.drawCentredString(w / 2, 85, "ARTHUR J. VILLANUEVA, PhD, CHA")
    c.setFont(F_SANS, 7.5)
    c.setFillColor(colors.HexColor("#64748B"))
    c.drawCentredString(w / 2, 72, "Chairman, Hospitality Operations Institute Examination Board")

    c.save()

    apply_torn_scan_defect(temp_clean_pdf, output_pdf, torn_box=(220, 175, 570, 215))
    if os.path.exists(safe_path(temp_clean_pdf)):
        os.remove(safe_path(temp_clean_pdf))
    return output_pdf


# ======================================================================
# GROUND TRUTH TXT FOR MARCUS ELIJAH NAVARRO
# ======================================================================
def build_marcus_ground_truth_txt(output_txt):
    content = f"""================================================================================
HOSPITALITY RESUME & SUPPORTING DOCUMENT BENCHMARK DATASET
GROUND TRUTH & VERIFICATION REFERENCE
================================================================================

DOCUMENT IDENTIFIER: Actual Info in the Resume of Marcus Elijah Navarro.txt
TARGET CANDIDATE: Marcus Elijah Navarro
CAREER DOMAIN: Hotel banquet operations, events, catering, function rooms, hospitality event services
TOTAL DATASET FILES: 13 Files (6 Resumes, 6 Supporting Documents, 1 Ground Truth)
DATE GENERATED: 2026-10-04

================================================================================
PART 1 — CANDIDATE MASTER PROFILE (STANDARDIZED CANONICAL RECORD)
================================================================================

[PERSONAL & CONTACT INFORMATION]
Full Name:             Marcus Elijah Navarro
Target Position:       Banquet Operations Supervisor
Contact Phone Number:  +63 917 552 8491
Email Address:         marcus.navarro.hospitality@outlook.ph
Residential Address:   Domestic Road, Pasay City, 1301 Metro Manila, Philippines
LinkedIn / Portfolio:  linkedin.com/in/marcus-navarro-hospitality

[PROFESSIONAL SUMMARY]
"Dedicated, results-driven Banquet Operations Supervisor with over 5 years of progressive hospitality leadership experience in luxury hotel convention centers and high-volume event venues. Proven expertise in Banquet Event Order (BEO) execution, floor logistics, VIP service protocols, team scheduling, and HACCP food safety standards. Adept at coordinating functions of up to 600 guests while ensuring seamless service delivery, minimal equipment breakage, and exceptional guest satisfaction."

[CORE COMPETENCIES & SKILLS]
1. Banquet Event Order (BEO) Floor Planning & Execution
2. Function Room Turnover & Service Floor Logistics
3. Staff Scheduling, Pre-Shift Briefings & Crew Leadership (up to 25 staff)
4. Formal Plated Service, Buffet Management & Banquet Bar Protocols
5. VIP Guest Relations & Protocol Handling
6. Breakage Mitigation & Chinaware/Glassware/Silverware (CGS) Par Levels
7. Food Safety, Sanitation, and HACCP Compliance
8. Banquet Billing, Beverage Requisition & Cost Control

[CHRONOLOGICAL WORK EXPERIENCE]

1. Employer:     Grand Crest Hotel & Suites
   Location:     Pasay City, Metro Manila
   Job Title:    Banquet Operations Supervisor
   Duration:     June 2022 – Present
   Tenure Type:  Current Employment
   Responsibilities & Achievements:
   • Supervise banquet operations for 3 grand ballrooms and 5 multi-function rooms catering up to 600 attendees per event.
   • Direct pre-shift briefings and manage a team of 18 full-time banquet servers, captains, and on-call catering staff.
   • Coordinate directly with culinary and sales teams to review BEO specifications, ensuring 99% on-time food rollout.
   • Reduced banquet equipment loss and chinaware breakage by 18% through standardized end-of-shift inventory protocols.

2. Employer:     Harborview Hotel & Convention Center
   Location:     Manila Bay, Manila
   Job Title:    Banquet Team Leader / Captain
   Duration:     January 2020 – May 2022
   Tenure Type:  Previous Employment
   Responsibilities & Achievements:
   • Led floor teams of 12 banquet attendants during corporate conventions, diplomatic receptions, and wedding galas.
   • Oversaw banquet table layouts, buffet staging, and silver service standards in compliance with luxury hotel guidelines.
   • Conducted daily equipment inspections and enforced food temperature logs in coordination with kitchen quality control.

3. Employer:     Bayfront Pavilion & Catering Services
   Location:     Parañaque City, Metro Manila
   Job Title:    Banquet Server / Function Attendant
   Duration:     July 2018 – December 2019
   Tenure Type:  Historical Employment
   Responsibilities & Achievements:
   • Delivered plated dinner service, replenished buffet lines, and executed room turnarounds for corporate and private events.
   • Maintained high standards of table grooming, glassware polishing, and rapid function hall clearing.

[EDUCATION & ACADEMIC CREDENTIALS]
Degree:         Bachelor of Science in Hospitality Management
Institution:    Pamantasan ng Lungsod ng Maynila (PLM) – Intramuros, Manila
Graduation:     2018

[TRAINING & PROFESSIONAL CERTIFICATIONS]
1. Title:   Certified Hospitality Supervisor (CHS)
   Issuer:  Hospitality Operations Institute / FHCB
   Date:    November 2021

2. Title:   Food Safety & Sanitation Standards (HACCP Level 2)
   Issuer:  Hospitality Training Council of the Philippines
   Date:    August 2022

3. Title:   Customer Service Excellence in Banquets and Events
   Issuer:  Philippine Hospitality Development Center
   Date:    March 2020


================================================================================
PART 2 — SIX RESUME GROUND TRUTHS
================================================================================

1. Marcus_Elijah_Navarro_Banquet_Operations_Supervisor.pdf
   - Format: PDF
   - Layout Design: Design A (Modern Executive Two-Column Navy Sidebar)
   - Target Position: Banquet Operations Supervisor
   - Career Emphasis: Banquet supervision, floor logistics, BEO execution, VIP service protocols.
   - Textual Content: Full Canonical Master Profile text preserved.

2. Marcus_Elijah_Navarro_Hotel_Events_Service_Coordinator.docx
   - Format: DOCX
   - Layout Design: Design B (Forest Green & Gold Top-Banner Layout)
   - Target Position: Hotel Events Service Coordinator
   - Career Emphasis: Event requirements, client coordination, function room setups, timeline enforcement.
   - Summary: Emphasizes bridging client expectations with hotel operational departments and BEO parameters.

3. Marcus_Elijah_Navarro_Banquet_Team_Leader.png
   - Format: PNG
   - Layout Design: Design C (Burgundy Top Header Banner & 2-Column Layout)
   - Target Position: Banquet Team Leader
   - Career Emphasis: Staff leadership, pre-shift briefings, service brigade management, on-the-floor coaching.
   - Summary: Action-oriented focus on motivating and coordinating banquet service crews during peak rushes.

4. Marcus_Elijah_Navarro_Catering_Service_Supervisor.jpg
   - Format: JPG
   - Layout Design: Design D (Warm Bronze & Metric Cards Minimalist Grid)
   - Target Position: Catering Service Supervisor
   - Career Emphasis: Catering operations, buffet presentation, beverage management, breakage control.
   - Summary: Emphasizes buffet logistics, equipment allocation, and prompt course rollout.

5. Marcus_Elijah_Navarro_Function_Operations_Officer_Blurred.png
   - Format: Blurred PNG (Uniform light blur applied for OCR testing)
   - Layout Design: Design E (Slate Teal Split Header & Card Layout)
   - Target Position: Function Operations Officer
   - Career Emphasis: Function room turnaround, audiovisual coordination, floor safety, equipment maintenance.
   - Text Degradation: Light uniform Gaussian blur across the entire image.

6. Marcus_Elijah_Navarro_Hotel_Events_Operations_Coordinator_Blurred.jpg
   - Format: Blurred JPG (Uniform light blur applied for OCR testing)
   - Layout Design: Design F (Royal Purple Classic Framed Layout)
   - Target Position: Hotel Events Operations Coordinator
   - Career Emphasis: Luxury event execution, guest journey, VIP protocol, high-volume event delivery.
   - Text Degradation: Light uniform defocus and compression simulation across the entire page.


================================================================================
PART 3 — SUPPORTING DOCUMENTS UNDERLYING TEXT & METADATA
================================================================================

1. Marcus_Elijah_Navarro_COE_01.pdf
   - Document Type: Certificate of Employment (COE)
   - Issuing Entity: Grand Crest Hotel & Suites
   - Recipient / Employee: Marcus Elijah Navarro
   - Certified Position: Banquet Operations Supervisor
   - Certified Period: June 2022 – Present
   - Extraction Mode: Digital Vector Text

2. Marcus_Elijah_Navarro_COE_02.pdf
   - Document Type: Certificate of Employment and Clearance
   - Issuing Entity: Harborview Hotel & Convention Center
   - Recipient / Employee: Marcus Elijah Navarro
   - Certified Position: Banquet Team Leader / Captain
   - Certified Period: January 15, 2021 – May 20, 2022 (CONTROLLED DISCREPANCY)
   - Extraction Mode: Digital Vector Text

3. Marcus_Elijah_Navarro_Diploma.pdf
   - Document Type: Academic University Diploma
   - Issuing Institution: Pamantasan ng Lungsod ng Maynila (PLM)
   - Graduate Name: Marcus Elijah Navarro
   - Conferred Degree: Bachelor of Science in Hospitality Management
   - Graduation Year: 2018 (April 18, 2018)
   - Extraction Mode: Digital Vector Text

4. Marcus_Elijah_Navarro_Training_Certificate_01.pdf
   - Document Type: Professional Training Certificate
   - Issuing Entity: Hospitality Training Council of the Philippines
   - Recipient Name: Marcus Elijah Navarro
   - Certified Course: Food Safety & Sanitation Standards (HACCP Level 2)
   - Certification Date: August 14, 2022
   - Extraction Mode: Digital Vector Text

5. Marcus_Elijah_Navarro_Training_Certificate_02.pdf
   - Document Type: Technical Workshop Certificate
   - Issuing Entity: Philippine Hospitality Development Center
   - Recipient Name: Marcus Elijah Navarro
   - Certified Course: Basic Table Setting and Napkin Folding Workshop (CONTROLLED DISCREPANCY)
   - Certification Date: March 18, 2020
   - Extraction Mode: Digital Vector Text

6. Marcus_Elijah_Navarro_Professional_Certification.pdf
   - Document Type: Professional Credential Certificate
   - Issuing Entity: Hospitality Operations Institute / FHCB
   - Recipient Name: [UNREADABLE / TORN / MISSING RECIPIENT NAME] (CONTROLLED UNABLE TO VERIFY)
   - Certified Title: Certified Hospitality Supervisor (CHS)
   - Issue Date: November 12, 2021
   - Extraction Mode: OCR-Test Document with Simulated Severe Physical Tear


================================================================================
PART 4 — RESUME ↔ SUPPORTING DOCUMENT VERIFICATION MATRIX
================================================================================

DOCUMENT: Marcus_Elijah_Navarro_COE_01.pdf
TYPE: Certificate of Employment
COMPARED RESUME CLAIM: Grand Crest Hotel & Suites | Banquet Operations Supervisor | June 2022 – Present
EXPECTED RESULT: VERIFIED
FIELDS:
  Employer   = MATCH (Grand Crest Hotel & Suites)
  Position   = MATCH (Banquet Operations Supervisor)
  Start Date = MATCH (June 2022)
  End Date   = MATCH (Present)
REASON: Exact match across candidate name, employer, role, and tenure.

--------------------------------------------------------------------------------

DOCUMENT: Marcus_Elijah_Navarro_COE_02.pdf
TYPE: Certificate of Employment and Clearance
COMPARED RESUME CLAIM: Harborview Hotel & Convention Center | Banquet Team Leader / Captain | January 2020 – May 2022
EXPECTED RESULT: DISCREPANCY_FOUND
FIELDS:
  Employer   = MATCH (Harborview Hotel & Convention Center)
  Position   = MATCH (Banquet Team Leader / Captain)
  Start Date = MISMATCH (Resume claims January 2020; COE certifies January 15, 2021)
  End Date   = MATCH (May 2022)
REASON: Official COE certifies employment starting January 2021, whereas resume claims start date of January 2020 (1-year discrepancy). Requires HR review.

--------------------------------------------------------------------------------

DOCUMENT: Marcus_Elijah_Navarro_Diploma.pdf
TYPE: Academic University Diploma
COMPARED RESUME CLAIM: Bachelor of Science in Hospitality Management | Pamantasan ng Lungsod ng Maynila | 2018
EXPECTED RESULT: VERIFIED
FIELDS:
  Institution = MATCH (Pamantasan ng Lungsod ng Maynila)
  Degree      = MATCH (Bachelor of Science in Hospitality Management)
  Year        = MATCH (2018)
REASON: Exact match across candidate name, institution, degree title, and graduation year.

--------------------------------------------------------------------------------

DOCUMENT: Marcus_Elijah_Navarro_Training_Certificate_01.pdf
TYPE: Training Certificate
COMPARED RESUME CLAIM: Food Safety & Sanitation Standards (HACCP Level 2) | Hospitality Training Council | August 2022
EXPECTED RESULT: VERIFIED
FIELDS:
  Course Title = MATCH (Food Safety & Sanitation Standards - HACCP Level 2)
  Issuer       = MATCH (Hospitality Training Council of the Philippines)
  Date         = MATCH (August 2022)
REASON: Exact match across candidate name, training title, issuer, and date.

--------------------------------------------------------------------------------

DOCUMENT: Marcus_Elijah_Navarro_Training_Certificate_02.pdf
TYPE: Technical Workshop Certificate
COMPARED RESUME CLAIM: Customer Service Excellence in Banquets and Events | Philippine Hospitality Development Center | March 2020
EXPECTED RESULT: DISCREPANCY_FOUND
FIELDS:
  Issuer       = MATCH (Philippine Hospitality Development Center)
  Date         = MATCH (March 2020)
  Course Title = MISMATCH (Resume claims 'Customer Service Excellence in Banquets and Events'; Certificate certifies 'Basic Table Setting and Napkin Folding Workshop')
REASON: Certificate course title does not match claimed training program. Requires HR review.

--------------------------------------------------------------------------------

DOCUMENT: Marcus_Elijah_Navarro_Professional_Certification.pdf
TYPE: Professional Certification
COMPARED RESUME CLAIM: Certified Hospitality Supervisor (CHS) | Hospitality Operations Institute / FHCB | November 2021
EXPECTED RESULT: UNABLE_TO_VERIFY
FIELDS:
  Recipient Name = UNREADABLE / MISSING (Severed due to torn scan artifact)
  Issuer         = MATCH (Hospitality Operations Institute / FHCB)
  Credential     = MATCH (Certified Hospitality Supervisor - CHS)
REASON: Recipient identity cannot be verified because the name portion of the credential document is severed/torn off.

--------------------------------------------------------------------------------

DOCUMENT: [DOCUMENT NOT SUBMITTED]
TYPE: Certificate of Employment (Historical)
COMPARED RESUME CLAIM: Bayfront Pavilion & Catering Services | Banquet Server / Function Attendant | July 2018 – December 2019
EXPECTED RESULT: PENDING
FIELDS:
  Evidence Status = PENDING_SUBMISSION
REASON: Supporting certificate of employment has not been provided by applicant for historical employment.
================================================================================
"""
    with open(safe_path(output_txt), 'w', encoding='utf-8') as f:
        f.write(content)
    return output_txt
